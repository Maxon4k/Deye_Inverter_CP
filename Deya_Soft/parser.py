import os
import re
import glob
import time
from datetime import datetime, timedelta

import xlrd
from openpyxl import load_workbook, Workbook

from config import DOWNLOAD_DIR


def parse_schedule_date(filename: str) -> str:
    """
    Витягує дату розкладу з імені файлу (наприклад DAM_22_04_2026.xlsx).
    Якщо не вдалось — повертає завтрашню дату.
    """
    basename = os.path.basename(filename)
    match    = re.search(r'(\d{2})[._](\d{2})[._](\d{4})', basename)
    if match:
        day, month, year = match.groups()
        result = datetime(int(year), int(month), int(day)).strftime("%Y-%m-%d")
        print(f"[DATE] {basename} → {result}")
        return result
    fallback = (datetime.now() + timedelta(days=1)).strftime("%Y-%m-%d")
    print(f"[DATE] Не вдалось розпізнати дату — використовую: {fallback}")
    return fallback


def convert_xls_to_xlsx(xls_path: str) -> str:
    """Конвертує .xls → .xlsx через xlrd + openpyxl."""
    xlsx_path = xls_path + "x"
    rb = None
    for ignore in (False, True):
        try:
            rb = xlrd.open_workbook(xls_path, ignore_workbook_corruption=ignore)
            if ignore:
                print(f"[CONVERT] ⚠️  Відкрито з ігноруванням пошкоджень")
            break
        except Exception as e:
            if not ignore:
                print(f"[CONVERT] Звичайне відкриття не вдалось: {e}")
            else:
                raise RuntimeError(f"Не вдалось відкрити .xls: {e}")
    rs = rb.sheet_by_index(0)
    wb = Workbook()
    ws = wb.active
    for row in range(rs.nrows):
        for col in range(rs.ncols):
            ws.cell(row=row + 1, column=col + 1).value = rs.cell_value(row, col)
    wb.save(xlsx_path)
    print(f"[CONVERT] → {os.path.basename(xlsx_path)}")
    return xlsx_path


def _clean_cell(value) -> float:
    if value is None:
        return 0.0
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def normalize_prices(raw_prices: list) -> list:
    """Нормалізує ціни до діапазону 1..100 (для запису в колонку C)."""
    min_p  = min(raw_prices)
    max_p  = max(raw_prices)
    spread = max_p - min_p
    if spread == 0:
        return [50.0] * len(raw_prices)
    return [(p - min_p) / spread * 99 + 1 for p in raw_prices]


def load_prices_from_file(file_path: str) -> tuple[list, str]:
    """
    Завантажує ціни з файлу .xls або .xlsx.

    Повертає (raw_prices, schedule_date).
    raw_prices — список з 24 значень грн/МВт (01:00..24:00).
    Також записує нормалізовані значення в колонку C файлу.
    """
    file_path = os.path.abspath(file_path)
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Файл не знайдено: {file_path}")

    # Видалити Zone.Identifier якщо є
    zone_id = file_path + ":Zone.Identifier"
    if os.path.exists(zone_id):
        try:
            os.remove(zone_id)
        except Exception:
            pass

    schedule_date = parse_schedule_date(file_path)
    print(f"[EXCEL] Обробка: {file_path}")

    if file_path.lower().endswith(".xls") and not file_path.lower().endswith(".xlsx"):
        file_path = convert_xls_to_xlsx(file_path)

    wb = load_workbook(file_path, data_only=True)
    ws = wb.active

    # РДН: row=2 → 01:00, row=25 → 24:00; ціна в колонці B
    raw_prices = [_clean_cell(ws.cell(row=i, column=2).value) for i in range(2, 26)]

    if len(raw_prices) < 24:
        print(f"[WARN] Знайдено лише {len(raw_prices)} рядків — очікується 24")

    # Нормалізовані значення → колонка C
    norm = normalize_prices(raw_prices)
    for idx, nv in enumerate(norm):
        ws.cell(row=idx + 2, column=3).value = round(nv, 2)
    wb.save(file_path)
    print(f"[EXCEL] Нормалізовані значення записано в колонку C")

    return raw_prices, schedule_date


def find_latest_file() -> str:
    """Шукає найновіший .xls/.xlsx файл у DOWNLOAD_DIR."""
    time.sleep(5)
    files = glob.glob(os.path.join(DOWNLOAD_DIR, "*.xls*"))
    if not files:
        raise FileNotFoundError(f"Файлів не знайдено в {DOWNLOAD_DIR}")
    return max(files, key=os.path.getctime)